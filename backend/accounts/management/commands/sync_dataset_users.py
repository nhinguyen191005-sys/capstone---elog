
from pathlib import Path

import openpyxl
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from accounts.models import Role, UserRole


EXPECTED_USERS = 12
EXPECTED_ASSIGNMENTS = 12
USER_SHEET = "02_Users"
ROLE_SHEET = "01_Roles"
USER_ROLE_SHEET = "03_User_Roles"



def read_sheet_records(workbook, sheet_name, required_columns):
    if sheet_name not in workbook.sheetnames:
        raise CommandError(f"Thiếu sheet bắt buộc: {sheet_name}")

    rows = workbook[sheet_name].iter_rows(values_only=True)
    header = None
    records = []

    # Dùng các cột bắt buộc để nhận diện header.
    # Sheet User_Roles không có cột id riêng.
    required = set(required_columns)

    for row_number, row in enumerate(rows, start=1):
        values = list(row)
        normalized = [
            str(value).strip() if value is not None else ""
            for value in values
        ]

        if header is None:
            if required.issubset(set(normalized)):
                header = normalized
                continue
            continue

        if not values or not any(v is not None for v in values):
            continue

        record = {
            name: values[index] if index < len(values) else None
            for index, name in enumerate(header)
        }

        # Với sheet có cột id, bỏ qua các dòng không có ID.
        # Với User_Roles, dùng user_id và role_id để nhận diện bản ghi.
        if "id" in record:
            if record["id"] is None:
                continue
        elif (
            record.get("user_id") is None
            or record.get("role_id") is None
        ):
            continue

        records.append(record)

    if header is None:
        raise CommandError(
            f"{sheet_name}: không tìm thấy hàng tiêu đề "
            f"chứa các cột {sorted(required)}."
        )

    return records



def required_text(value, field, sheet, source_id):
    if value is None or not str(value).strip():
        raise CommandError(
            f"{sheet}, ID {source_id}: thiếu {field}."
        )
    return str(value).strip()


class Command(BaseCommand):
    help = (
        "Kiểm tra hoặc đồng bộ 12 tài khoản demo và phân công vai trò "
        "từ CoLogi_Dataset.xlsx."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            default="../data/CoLogi_Dataset.xlsx",
            help="Đường dẫn workbook, tính từ thư mục backend.",
        )
        parser.add_argument(
            "--apply",
            action="store_true",
            help="Ghi dữ liệu sau khi xác thực toàn bộ dữ liệu nguồn.",
        )

    def handle(self, *args, **options):
        file_path = Path(options["file"]).resolve()

        if not file_path.is_file():
            raise CommandError(f"Không tìm thấy file: {file_path}")

        workbook = openpyxl.load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        try:
            users = read_sheet_records(
                workbook,
                USER_SHEET,
                ["id", "full_name", "email", "is_active"],
            )
            roles = read_sheet_records(
                workbook,
                ROLE_SHEET,
                ["id", "role_code", "role_name", "is_active"],
            )
            assignments = read_sheet_records(
                workbook,
                USER_ROLE_SHEET,
                ["user_id", "role_id", "assigned_by"],
            )

            if len(users) != EXPECTED_USERS:
                raise CommandError(
                    f"{USER_SHEET}: mong đợi {EXPECTED_USERS} "
                    f"tài khoản, đọc được {len(users)}."
                )

            if len(assignments) != EXPECTED_ASSIGNMENTS:
                raise CommandError(
                    f"{USER_ROLE_SHEET}: mong đợi "
                    f"{EXPECTED_ASSIGNMENTS} phân công, "
                    f"đọc được {len(assignments)}."
                )

            # Xác thực dữ liệu người dùng và tạo bản đồ ID nguồn.
            user_by_source_id = {}
            usernames_seen = set()
            emails_seen = set()

            for record in users:
                source_id = record["id"]

                if source_id in user_by_source_id:
                    raise CommandError(
                        f"{USER_SHEET}: trùng ID {source_id}."
                    )

                email = required_text(
                    record["email"], "email", USER_SHEET, source_id
                ).lower()

                full_name = required_text(
                    record["full_name"],
                    "full_name",
                    USER_SHEET,
                    source_id,
                )

                # Email trong workbook có dạng userNN@cologi.demo.
                username = email.split("@", 1)[0]

                if "@" not in email:
                    raise CommandError(
                        f"{USER_SHEET}, ID {source_id}: email không hợp lệ."
                    )

                if username in usernames_seen or email in emails_seen:
                    raise CommandError(
                        f"{USER_SHEET}: username/email bị trùng "
                        f"ở tài khoản ID {source_id}."
                    )

                usernames_seen.add(username)
                emails_seen.add(email)

                if not isinstance(record["is_active"], bool):
                    raise CommandError(
                        f"{USER_SHEET}, ID {source_id}: "
                        "is_active không phải giá trị Boolean."
                    )

                user_by_source_id[source_id] = {
                    "source_id": source_id,
                    "username": username,
                    "email": email,
                    "full_name": full_name,
                    "is_active": record["is_active"],
                }

            # Xác thực các vai trò nguồn.
            role_by_source_id = {}
            role_codes_seen = set()

            for record in roles:
                source_id = record["id"]
                role_code = required_text(
                    record["role_code"],
                    "role_code",
                    ROLE_SHEET,
                    source_id,
                )

                if source_id in role_by_source_id:
                    raise CommandError(
                        f"{ROLE_SHEET}: trùng ID {source_id}."
                    )

                if role_code in role_codes_seen:
                    raise CommandError(
                        f"{ROLE_SHEET}: trùng role_code {role_code}."
                    )

                role_codes_seen.add(role_code)
                role_by_source_id[source_id] = role_code

            # Xác thực phân công; mỗi tài khoản chỉ có một vai trò
            # theo dữ liệu hiện tại của workbook.
            prepared_assignments = []
            assignment_pairs = set()

            for record in assignments:
                user_source_id = record["user_id"]
                role_source_id = record["role_id"]
                assigned_by_source_id = record["assigned_by"]

                if user_source_id not in user_by_source_id:
                    raise CommandError(
                        f"{USER_ROLE_SHEET}: user_id "
                        f"{user_source_id} không tồn tại trong {USER_SHEET}."
                    )

                if role_source_id not in role_by_source_id:
                    raise CommandError(
                        f"{USER_ROLE_SHEET}: role_id "
                        f"{role_source_id} không tồn tại trong {ROLE_SHEET}."
                    )

                if assigned_by_source_id not in user_by_source_id:
                    raise CommandError(
                        f"{USER_ROLE_SHEET}: assigned_by "
                        f"{assigned_by_source_id} không tồn tại trong "
                        f"{USER_SHEET}."
                    )

                pair = (user_source_id, role_source_id)

                if pair in assignment_pairs:
                    raise CommandError(
                        f"{USER_ROLE_SHEET}: phân công trùng {pair}."
                    )

                assignment_pairs.add(pair)
                prepared_assignments.append(
                    {
                        "user_source_id": user_source_id,
                        "role_code": role_by_source_id[role_source_id],
                        "assigned_by_source_id": assigned_by_source_id,
                    }
                )

            if len(assignment_pairs) != EXPECTED_ASSIGNMENTS:
                raise CommandError(
                    "Số lượng phân công vai trò không đúng kỳ vọng."
                )

            User = get_user_model()

            # Đảm bảo cả năm role_code trong workbook đều có trong DB.
            db_roles = {
                role.role_code: role
                for role in Role.objects.filter(
                    role_code__in=role_codes_seen
                )
            }

            missing_roles = role_codes_seen - set(db_roles)
            if missing_roles:
                raise CommandError(
                    f"Database thiếu vai trò: {sorted(missing_roles)}."
                )

            inactive_roles = sorted(
                code for code in role_codes_seen
                if not db_roles[code].is_active
            )
            if inactive_roles:
                raise CommandError(
                    f"Vai trò đang bị vô hiệu hóa: {inactive_roles}."
                )

            # Không chiếm username/email của tài khoản hiện hữu.
            existing_usernames = set(
                User.objects.filter(
                    username__in=usernames_seen
                ).values_list("username", flat=True)
            )
            existing_emails = set(
                User.objects.filter(
                    email__in=emails_seen
                ).values_list("email", flat=True)
            )

            if existing_usernames or existing_emails:
                raise CommandError(
                    "Có xung đột với tài khoản hiện hữu. "
                    f"Username trùng: {sorted(existing_usernames)}; "
                    f"email trùng: {sorted(existing_emails)}. "
                    "Không có dữ liệu nào được ghi."
                )

            self.stdout.write(f"Workbook: {file_path}")
            self.stdout.write(f"Tài khoản hợp lệ: {len(user_by_source_id)}")
            self.stdout.write(
                f"Phân công vai trò hợp lệ: {len(prepared_assignments)}"
            )
            self.stdout.write(
                f"Vai trò đã đối chiếu: {len(role_codes_seen)}"
            )
            self.stdout.write(
                "Tài khoản hiện có và các phân công hiện tại sẽ được giữ nguyên."
            )
            self.stdout.write(
                "Mật khẩu mới sẽ ở trạng thái chưa thiết lập."
            )

            if not options["apply"]:
                self.stdout.write(
                    self.style.SUCCESS(
                        "DRY RUN hoàn tất. Database chưa bị thay đổi. "
                        "Chạy lại với --apply khi đã sẵn sàng."
                    )
                )
                return

            with transaction.atomic():
                # Kiểm tra lại xung đột trong transaction trước khi ghi.
                if User.objects.filter(
                    username__in=usernames_seen
                ).exists():
                    raise CommandError(
                        "Username đã xuất hiện trước khi ghi; đã hủy."
                    )

                if User.objects.filter(email__in=emails_seen).exists():
                    raise CommandError(
                        "Email đã xuất hiện trước khi ghi; đã hủy."
                    )

                # Tạo tài khoản bằng API Django; không dùng password_hash
                # placeholder của workbook.
                created_users = {}

                for source_id, data in user_by_source_id.items():
                    user = User(
                        username=data["username"],
                        email=data["email"],
                        first_name=data["full_name"],
                        is_active=data["is_active"],
                        is_staff=False,
                        is_superuser=False,
                    )
                    user.set_unusable_password()
                    user.save()
                    created_users[source_id] = user

                # Tạo phân công sau khi đã có bản đồ ID nguồn -> user DB.
                for item in prepared_assignments:
                    user = created_users[item["user_source_id"]]
                    role = db_roles[item["role_code"]]
                    assigned_by = created_users[
                        item["assigned_by_source_id"]
                    ]

                    UserRole.objects.create(
                        user=user,
                        role=role,
                        assigned_by=assigned_by,
                    )

                # Kiểm tra kết quả trước khi commit.
                created_count = sum(
                    User.objects.filter(
                        pk=user.pk
                    ).exists()
                    for user in created_users.values()
                )

                assignment_count = UserRole.objects.filter(
                    user__in=list(created_users.values())
                ).count()

                if created_count != EXPECTED_USERS:
                    raise CommandError(
                        f"Đã tạo {created_count}/{EXPECTED_USERS} tài khoản."
                    )

                if assignment_count != EXPECTED_ASSIGNMENTS:
                    raise CommandError(
                        "Số phân công sau ghi không đúng kỳ vọng: "
                        f"{assignment_count}/{EXPECTED_ASSIGNMENTS}."
                    )

            self.stdout.write(
                self.style.SUCCESS(
                    f"Hoàn tất: tạo {EXPECTED_USERS} tài khoản và "
                    f"{EXPECTED_ASSIGNMENTS} phân công vai trò."
                )
            )
            self.stdout.write(
                "Các tài khoản mới chưa có mật khẩu. "
                "Hãy đặt mật khẩu qua Django trước khi thử đăng nhập."
            )

        finally:
            workbook.close()
