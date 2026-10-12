import json
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl
from django.contrib.auth import get_user_model
from django.core.management.base import BaseCommand, CommandError
from django.db import connection, transaction
from django.db.models import Max
from django.core.management.color import no_style

from masterdata.models import (
    RuleSet,
    BusinessRule,
    PlanningParameter,
    CostProfile,
)


EXPECTED_COUNTS = {
    "11_Rule_Sets": 3,
    "12_Business_Rules": 84,
    "13_Planning_Params": 10,
    "14_Cost_Profiles": 5,
}

TARGET_MODELS = (RuleSet, BusinessRule, PlanningParameter, CostProfile)
SOURCE_TZ = ZoneInfo("Asia/Ho_Chi_Minh")


def normalize_id(value, label):
    if isinstance(value, bool) or value is None:
        raise CommandError(f"{label}: ID không hợp lệ: {value!r}")
    try:
        number = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise CommandError(f"{label}: ID không hợp lệ: {value!r}")
    if not number.is_finite() or number != number.to_integral_value():
        raise CommandError(f"{label}: ID phải là số nguyên: {value!r}")
    result = int(number)
    if result <= 0:
        raise CommandError(f"{label}: ID phải lớn hơn 0.")
    return result


def required_text(value, label):
    if value is None or not str(value).strip():
        raise CommandError(f"Thiếu trường bắt buộc: {label}.")
    return str(value).strip()


def as_bool(value, label):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        normalized = value.strip().lower()
        if normalized in {"true", "1", "yes"}:
            return True
        if normalized in {"false", "0", "no"}:
            return False
    raise CommandError(f"{label}: giá trị boolean không hợp lệ: {value!r}")


def as_datetime(value, label, nullable=False):
    if value is None and nullable:
        return None
    if not isinstance(value, datetime):
        if isinstance(value, date):
            value = datetime.combine(value, time.min)
        else:
            raise CommandError(f"{label}: ngày giờ không hợp lệ: {value!r}")
    if value.tzinfo is None:
        return value.replace(tzinfo=SOURCE_TZ)
    return value.astimezone(SOURCE_TZ)


def as_decimal(value, label, nullable=False):
    if value is None and nullable:
        return None
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise CommandError(f"{label}: số không hợp lệ: {value!r}")
    if not result.is_finite():
        raise CommandError(f"{label}: số không hữu hạn.")
    return result


def parse_json(value, label):
    if not isinstance(value, str):
        if isinstance(value, (dict, list, int, float, bool)) or value is None:
            return value
        raise CommandError(f"{label}: JSON không hợp lệ.")
    try:
        return json.loads(value)
    except json.JSONDecodeError as exc:
        raise CommandError(f"{label}: JSON không hợp lệ: {exc.msg}") from exc


def read_records(workbook, sheet_name, required_columns):
    if sheet_name not in workbook.sheetnames:
        raise CommandError(f"Thiếu sheet bắt buộc: {sheet_name}")

    header = None
    records = []
    required = set(required_columns)

    for row_number, row in enumerate(
        workbook[sheet_name].iter_rows(values_only=True), start=1
    ):
        values = list(row)
        normalized = [
            str(value).strip() if value is not None else ""
            for value in values
        ]

        if header is None:
            if required.issubset(set(normalized)):
                header = normalized
            continue

        if not any(value is not None for value in values):
            continue

        record = {
            name: values[index] if index < len(values) else None
            for index, name in enumerate(header)
            if name
        }
        if record.get("id") is not None:
            records.append((row_number, record))

    if header is None:
        raise CommandError(
            f"{sheet_name}: không tìm thấy header có các cột "
            f"{sorted(required)}."
        )
    return records


class Command(BaseCommand):
    help = "Kiểm tra hoặc import 4 bảng cấu hình từ CoLogi_Dataset.xlsx."

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            default="../data/CoLogi_Dataset.xlsx",
            help="Đường dẫn workbook, tính từ thư mục backend.",
        )
        parser.add_argument(
            "--apply",
            action="store_true",
            help="Ghi dữ liệu sau khi kiểm tra toàn bộ nguồn.",
        )

    def handle(self, *args, **options):
        file_path = Path(options["file"]).resolve()
        if not file_path.is_file():
            raise CommandError(f"Không tìm thấy file: {file_path}")

        workbook = openpyxl.load_workbook(
            file_path, read_only=True, data_only=True
        )

        try:
            rule_set_rows = read_records(
                workbook,
                "11_Rule_Sets",
                [
                    "id", "rule_set_code", "rule_set_name", "version_no",
                    "status", "effective_from", "created_by",
                ],
            )
            business_rule_rows = read_records(
                workbook,
                "12_Business_Rules",
                [
                    "id", "rule_set_id", "rule_code", "module_code",
                    "rule_name", "description", "rule_definition",
                    "priority", "is_active",
                ],
            )
            parameter_rows = read_records(
                workbook,
                "13_Planning_Params",
                [
                    "id", "parameter_key", "parameter_value", "data_type",
                    "description", "updated_by", "updated_at",
                ],
            )
            cost_profile_rows = read_records(
                workbook,
                "14_Cost_Profiles",
                [
                    "id", "profile_code", "profile_name", "currency_code",
                    "fixed_transport_cost", "cost_per_km",
                    "additional_handling_cost", "is_active",
                ],
            )

            rows_by_sheet = {
                "11_Rule_Sets": rule_set_rows,
                "12_Business_Rules": business_rule_rows,
                "13_Planning_Params": parameter_rows,
                "14_Cost_Profiles": cost_profile_rows,
            }
            for sheet, expected in EXPECTED_COUNTS.items():
                actual = len(rows_by_sheet[sheet])
                if actual != expected:
                    raise CommandError(
                        f"{sheet}: mong đợi {expected}, đọc được {actual}."
                    )

            # Ánh xạ ID người dùng nguồn sang email, rồi sang User trong DB.
            user_rows = read_records(
                workbook,
                "02_Users",
                ["id", "email", "full_name", "is_active"],
            )
            User = get_user_model()
            source_user_emails = {}
            for row_number, record in user_rows:
                label = f"02_Users, dòng {row_number}"
                source_id = normalize_id(record["id"], label)
                email = required_text(record["email"], f"{label}, email").lower()
                if source_id in source_user_emails:
                    raise CommandError(f"{label}: ID người dùng nguồn bị trùng.")
                source_user_emails[source_id] = email

            db_users = {
                user.email.lower(): user
                for user in User.objects.filter(
                    email__in=set(source_user_emails.values())
                )
            }
            missing_emails = set(source_user_emails.values()) - set(db_users)
            if missing_emails:
                raise CommandError(
                    f"Không tìm thấy tài khoản Django: {sorted(missing_emails)}"
                )

            user_map = {
                source_id: db_users[email]
                for source_id, email in source_user_emails.items()
            }

            # Chuẩn bị RuleSet.
            prepared_rule_sets = []
            source_rule_set_ids = set()
            rule_set_codes = set()
            rule_set_versions = set()
            valid_rule_set_statuses = {item.value for item in RuleSet.Status}

            for row_number, record in rule_set_rows:
                label = f"11_Rule_Sets, dòng {row_number}"
                source_id = normalize_id(record["id"], label)
                if source_id in source_rule_set_ids:
                    raise CommandError(f"{label}: ID bị trùng.")
                source_rule_set_ids.add(source_id)

                code = required_text(record["rule_set_code"], f"{label}, code")
                name = required_text(record["rule_set_name"], f"{label}, name")
                version = normalize_id(record["version_no"], f"{label}, version")
                status = required_text(record["status"], f"{label}, status")
                creator_id = normalize_id(record["created_by"], f"{label}, created_by")

                if status not in valid_rule_set_statuses:
                    raise CommandError(f"{label}: status không hợp lệ: {status}")
                if creator_id not in user_map:
                    raise CommandError(
                        f"{label}: created_by={creator_id} không có trong 02_Users."
                    )

                unique_key = (code, version)
                if unique_key in rule_set_versions:
                    raise CommandError(f"{label}: mã và phiên bản bị trùng.")
                rule_set_versions.add(unique_key)

                prepared_rule_sets.append({
                    "source_id": source_id,
                    "object": RuleSet(
                        id=source_id,
                        rule_set_code=code,
                        rule_set_name=name,
                        version_no=version,
                        status=status,
                        effective_from=as_datetime(
                            record["effective_from"],
                            f"{label}, effective_from",
                            nullable=True,
                        ),
                        created_by=user_map[creator_id],
                    ),
                })

            # Chuẩn bị BusinessRule.
            prepared_rules = []
            rule_ids = set()
            rule_codes_by_set = set()
            valid_modules = {item.value for item in BusinessRule.ModuleCode}

            for row_number, record in business_rule_rows:
                label = f"12_Business_Rules, dòng {row_number}"
                source_id = normalize_id(record["id"], label)
                if source_id in rule_ids:
                    raise CommandError(f"{label}: ID bị trùng.")
                rule_ids.add(source_id)

                rule_set_source_id = normalize_id(record["rule_set_id"], label)
                if rule_set_source_id not in source_rule_set_ids:
                    raise CommandError(
                        f"{label}: rule_set_id={rule_set_source_id} không tồn tại."
                    )

                rule_code = required_text(record["rule_code"], f"{label}, rule_code")
                module_code = required_text(record["module_code"], f"{label}, module_code")
                if module_code not in valid_modules:
                    raise CommandError(
                        f"{label}: module_code không hợp lệ: {module_code}"
                    )

                unique_key = (rule_set_source_id, rule_code)
                if unique_key in rule_codes_by_set:
                    raise CommandError(f"{label}: rule_code trùng trong RuleSet.")
                rule_codes_by_set.add(unique_key)

                prepared_rules.append({
                    "source_id": source_id,
                    "object": BusinessRule(
                        id=source_id,
                        rule_set_id=rule_set_source_id,
                        rule_code=rule_code,
                        module_code=module_code,
                        rule_name=required_text(record["rule_name"], f"{label}, rule_name"),
                        description=record.get("description"),
                        rule_definition=parse_json(
                            record["rule_definition"],
                            f"{label}, rule_definition",
                        ),
                        priority=normalize_id(record["priority"], f"{label}, priority"),
                        is_active=as_bool(record["is_active"], f"{label}, is_active"),
                    ),
                })

            # Chuẩn bị PlanningParameter.
            prepared_parameters = []
            parameter_ids = set()
            parameter_keys = set()
            for row_number, record in parameter_rows:
                label = f"13_Planning_Params, dòng {row_number}"
                source_id = normalize_id(record["id"], label)
                if source_id in parameter_ids:
                    raise CommandError(f"{label}: ID bị trùng.")
                parameter_ids.add(source_id)

                key = required_text(record["parameter_key"], f"{label}, parameter_key")
                if key in parameter_keys:
                    raise CommandError(f"{label}: parameter_key bị trùng.")
                parameter_keys.add(key)

                updated_by_source = record.get("updated_by")
                if updated_by_source is None:
                    updated_by = None
                else:
                    updated_by_id = normalize_id(updated_by_source, f"{label}, updated_by")
                    if updated_by_id not in user_map:
                        raise CommandError(
                            f"{label}: updated_by={updated_by_id} không có trong 02_Users."
                        )
                    updated_by = user_map[updated_by_id]

                prepared_parameters.append({
                    "source_id": source_id,
                    "object": PlanningParameter(
                        id=source_id,
                        parameter_key=key,
                        parameter_value=parse_json(
                            record["parameter_value"],
                            f"{label}, parameter_value",
                        ),
                        data_type=required_text(record["data_type"], f"{label}, data_type"),
                        description=record.get("description"),
                        updated_by=updated_by,
                        updated_at=as_datetime(record["updated_at"], f"{label}, updated_at"),
                    ),
                })

            # Chuẩn bị CostProfile.
            prepared_cost_profiles = []
            cost_profile_ids = set()
            profile_codes = set()
            for row_number, record in cost_profile_rows:
                label = f"14_Cost_Profiles, dòng {row_number}"
                source_id = normalize_id(record["id"], label)
                if source_id in cost_profile_ids:
                    raise CommandError(f"{label}: ID bị trùng.")
                cost_profile_ids.add(source_id)

                code = required_text(record["profile_code"], f"{label}, profile_code")
                if code in profile_codes:
                    raise CommandError(f"{label}: profile_code bị trùng.")
                profile_codes.add(code)

                fixed = as_decimal(record["fixed_transport_cost"], f"{label}, fixed_transport_cost")
                per_km = as_decimal(record["cost_per_km"], f"{label}, cost_per_km", nullable=True)
                handling = as_decimal(record["additional_handling_cost"], f"{label}, additional_handling_cost")
                if fixed < 0 or (per_km is not None and per_km < 0) or handling < 0:
                    raise CommandError(f"{label}: chi phí không được âm.")

                prepared_cost_profiles.append({
                    "source_id": source_id,
                    "object": CostProfile(
                        id=source_id,
                        profile_code=code,
                        profile_name=required_text(record["profile_name"], f"{label}, profile_name"),
                        currency_code=required_text(record["currency_code"], f"{label}, currency_code"),
                        fixed_transport_cost=fixed,
                        cost_per_km=per_km,
                        additional_handling_cost=handling,
                        is_active=as_bool(record["is_active"], f"{label}, is_active"),
                    ),
                })

            # Chặn các khóa tự nhiên đã tồn tại và kiểm tra bảng đích.
            if RuleSet.objects.exists() or BusinessRule.objects.exists() or PlanningParameter.objects.exists() or CostProfile.objects.exists():
                raise CommandError(
                    "Một hoặc nhiều bảng cấu hình đã có dữ liệu. "
                    "Importer chỉ hỗ trợ nhập khi cả bốn bảng đích đều trống."
                )

            existing_user_ids = set(user_map.values())
            if not existing_user_ids:
                raise CommandError("Không ánh xạ được tài khoản người dùng.")

            print_summary = [
                ("RuleSet", prepared_rule_sets),
                ("BusinessRule", prepared_rules),
                ("PlanningParameter", prepared_parameters),
                ("CostProfile", prepared_cost_profiles),
            ]
            self.stdout.write(f"Workbook: {file_path}")
            for name, items in print_summary:
                self.stdout.write(f"{name}: {len(items)} bản ghi hợp lệ")
            self.stdout.write("Tài khoản nguồn đã ánh xạ qua email.")

            if not options["apply"]:
                self.stdout.write(self.style.SUCCESS(
                    "DRY RUN hoàn tất. Database chưa bị thay đổi. "
                    "Chạy lại với --apply khi đã sẵn sàng."
                ))
                return

            with transaction.atomic():
                for model in TARGET_MODELS:
                    if model.objects.exists():
                        raise CommandError(
                            f"{model.__name__} không còn trống; hủy import."
                        )

                # Tạo theo thứ tự cha trước, con sau.
                RuleSet.objects.bulk_create(
                    [item["object"] for item in prepared_rule_sets],
                    batch_size=100,
                )
                BusinessRule.objects.bulk_create(
                    [item["object"] for item in prepared_rules],
                    batch_size=200,
                )
                PlanningParameter.objects.bulk_create(
                    [item["object"] for item in prepared_parameters],
                    batch_size=100,
                )
                CostProfile.objects.bulk_create(
                    [item["object"] for item in prepared_cost_profiles],
                    batch_size=100,
                )

                expected_by_model = {
                    RuleSet: EXPECTED_COUNTS["11_Rule_Sets"],
                    BusinessRule: EXPECTED_COUNTS["12_Business_Rules"],
                    PlanningParameter: EXPECTED_COUNTS["13_Planning_Params"],
                    CostProfile: EXPECTED_COUNTS["14_Cost_Profiles"],
                }
                for model, expected in expected_by_model.items():
                    actual = model.objects.count()
                    if actual != expected:
                        raise CommandError(
                            f"{model.__name__}: mong đợi {expected}, thực tế {actual}; rollback."
                        )

                # Đồng bộ sequence cho các ID được nhập thủ công.
                sequence_sql = connection.ops.sequence_reset_sql(
                    no_style(), list(TARGET_MODELS)
                )
                with connection.cursor() as cursor:
                    for sql in sequence_sql:
                        cursor.execute(sql)

            self.stdout.write(self.style.SUCCESS(
                "Import hoàn tất: 3 RuleSet, 84 BusinessRule, "
                "10 PlanningParameter, 5 CostProfile."
            ))
        finally:
            workbook.close()
