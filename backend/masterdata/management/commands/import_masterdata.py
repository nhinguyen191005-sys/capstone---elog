from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from masterdata.models import (
    CargoCategory,
    Hub,
    Outlet,
    Supplier,
    VehicleType,
)


SHEET_CONFIG = [
    {
        "sheet": "04_Suppliers",
        "model": Supplier,
        "key": "supplier_code",
        "fields": [
            "supplier_code",
            "supplier_name",
            "contact_name",
            "phone",
            "email",
            "address",
            "is_active",
            "created_at",
            "updated_at",
        ],
        "expected_count": 10,
    },
    {
        "sheet": "05_Outlets",
        "model": Outlet,
        "key": "outlet_code",
        "fields": [
            "outlet_code",
            "outlet_name",
            "address",
            "contact_name",
            "phone",
            "is_active",
            "created_at",
            "updated_at",
        ],
        "expected_count": 12,
    },
    {
        "sheet": "06_Hubs",
        "model": Hub,
        "key": "hub_code",
        "fields": [
            "hub_code",
            "hub_name",
            "address",
            "is_active",
            "created_at",
        ],
        "expected_count": 3,
    },
    {
        "sheet": "07_Cargo_Categories",
        "model": CargoCategory,
        "key": "category_code",
        "fields": [
            "category_code",
            "category_name",
            "description",
            "is_active",
        ],
        "expected_count": 7,
    },
    {
        "sheet": "09_Vehicle_Types",
        "model": VehicleType,
        "key": "vehicle_code",
        "fields": [
            "vehicle_code",
            "vehicle_name",
            "capacity_kg",
            "description",
            "is_active",
        ],
        "expected_count": 7,
    },
]


def normalize_value(field_name, value):
    if value is None:
        return None

    if field_name == "capacity_kg":
        try:
            return Decimal(str(value))
        except InvalidOperation as exc:
            raise CommandError(
                f"capacity_kg không hợp lệ: {value!r}"
            ) from exc

    return value


class Command(BaseCommand):
    help = "Import 5 bảng master data từ CoLogi_Dataset.xlsx."

    def add_arguments(self, parser):
        parser.add_argument(
            "--file",
            default="../data/CoLogi_Dataset.xlsx",
            help="Đường dẫn đến file Excel, tính từ thư mục backend.",
        )
        parser.add_argument(
            "--dry-run",
            action="store_true",
            help="Chỉ kiểm tra dữ liệu, không ghi vào database.",
        )

    def handle(self, *args, **options):
        file_path = Path(options["file"]).resolve()

        if not file_path.is_file():
            raise CommandError(f"Không tìm thấy file: {file_path}")

        workbook = openpyxl.load_workbook(
            file_path,
            read_only=True,
            data_only=True,
        )

        try:
            prepared = []

            # Đọc và xác thực toàn bộ dữ liệu trước khi ghi.
            for config in SHEET_CONFIG:
                sheet_name = config["sheet"]

                if sheet_name not in workbook.sheetnames:
                    raise CommandError(
                        f"Thiếu sheet bắt buộc: {sheet_name}"
                    )

                rows = list(
                    workbook[sheet_name].iter_rows(values_only=True)
                )

                header_idx = next(
                    (
                        i
                        for i, row in enumerate(rows)
                        if row
                        and str(row[0]).strip().lower() == "id"
                    ),
                    None,
                )

                if header_idx is None:
                    raise CommandError(
                        f"{sheet_name}: không tìm thấy dòng tiêu đề."
                    )

                header = [
                    str(value).strip() if value is not None else ""
                    for value in rows[header_idx]
                ]

                missing = [
                    field
                    for field in config["fields"]
                    if field not in header
                ]

                if missing:
                    raise CommandError(
                        f"{sheet_name}: thiếu cột {missing}"
                    )

                key_idx = header.index(config["key"])
                seen_keys = set()
                records = []

                for row_number, row in enumerate(
                    rows[header_idx + 1:],
                    start=header_idx + 2,
                ):
                    if not row or row[0] is None:
                        continue

                    record = {}
                    for field_name in config["fields"]:
                        index = header.index(field_name)
                        value = row[index] if index < len(row) else None
                        record[field_name] = normalize_value(
                            field_name, value
                        )

                    business_key = record[config["key"]]

                    if not business_key:
                        raise CommandError(
                            f"{sheet_name}, dòng {row_number}: "
                            f"thiếu {config['key']}."
                        )

                    if business_key in seen_keys:
                        raise CommandError(
                            f"{sheet_name}: trùng "
                            f"{config['key']}={business_key!r}."
                        )

                    seen_keys.add(business_key)
                    records.append(record)

                if len(records) != config["expected_count"]:
                    raise CommandError(
                        f"{sheet_name}: dự kiến "
                        f"{config['expected_count']} bản ghi, "
                        f"nhưng đọc được {len(records)}."
                    )

                prepared.append((config, records))

            if options["dry_run"]:
                for config, records in prepared:
                    self.stdout.write(
                        f"[DRY RUN] {config['sheet']}: "
                        f"{len(records)} bản ghi hợp lệ."
                    )
                self.stdout.write(
                    self.style.SUCCESS(
                        "Kiểm tra hoàn tất; database chưa bị thay đổi."
                    )
                )
                return

            # Không ghi đè bất kỳ bảng danh mục nào đã có dữ liệu.
            with transaction.atomic():
                for config, records in prepared:
                    model = config["model"]

                    if model.objects.exists():
                        raise CommandError(
                            f"{model.__name__} không còn trống. "
                            "Đã hủy import để tránh ghi đè dữ liệu."
                        )

                for config, records in prepared:
                    model = config["model"]
                    model.objects.bulk_create(
                        [model(**record) for record in records]
                    )

                # Xác minh số lượng trước khi transaction commit.
                for config, records in prepared:
                    actual = config["model"].objects.count()
                    expected = config["expected_count"]

                    if actual != expected:
                        raise CommandError(
                            f"{config['model'].__name__}: "
                            f"mong đợi {expected}, thực tế {actual}."
                        )

            for config, records in prepared:
                self.stdout.write(
                    f"{config['model'].__name__}: "
                    f"đã import {len(records)} bản ghi."
                )

            self.stdout.write(
                self.style.SUCCESS(
                    "Import 5 bảng danh mục hoàn tất."
                )
            )
        finally:
            workbook.close()
